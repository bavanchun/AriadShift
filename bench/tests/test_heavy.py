"""Independent heavy-fixture truth, visual degradation and public-report checks."""

import json
from pathlib import Path
import tomllib
from zipfile import ZipFile

from lxml import etree
import numpy as np
from openpyxl import load_workbook
import pytest

from ariad_bench.canonical import normalize
from ariad_bench.report import render
from ariad_bench.visual import page_ssim, pdf_ssim, raster_pages
from ariad_fixture_gen.heavy import generate
from ariad_fixture_gen.heavy_office import generate as generate_office

ROOT = Path(__file__).resolve().parents[2]


def test_heavy_fixture_generation_is_reproducible_and_manifested():
    first = generate(ROOT) | generate_office(ROOT)
    second = generate(ROOT) | generate_office(ROOT)
    assert first == second
    records = tomllib.loads((ROOT / "fixtures/manifest.toml").read_text())["fixture"]
    declared = {
        path
        for f in records
        if "heavy-bench" in f.get("tags", [])
        for path in [f["path"], *(c["path"] for c in f["companions"])]
    }
    assert set(first) == declared
    for path, content in first.items():
        assert (ROOT / path).read_bytes() == content, path


@pytest.mark.parametrize("stem", ["en-map", "vi-garden"])
def test_pdf_and_office_truth_share_authored_content(stem):
    base = ROOT / "fixtures/heavy" / stem
    truth = normalize(base.with_suffix(".pdf.truth.json"))
    assert len(truth.headings) == 1
    assert len(truth.tables) == 1
    assert len(raster_pages(base.with_suffix(".pdf"))) == 1
    assert pdf_ssim(base.with_suffix(".pdf"), base.parent / f"{stem}-scan.pdf") < 1
    with ZipFile(base.with_suffix(".docx")) as archive:
        tree = etree.fromstring(archive.read("word/document.xml"))
        body_text = " ".join(tree.xpath("//*[local-name()='t']/text()"))
        assert " ".join(truth.text.split()) == " ".join(body_text.split())
    with ZipFile(base.with_suffix(".pptx")) as archive:
        tree = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        actual = " ".join(tree.xpath("//*[local-name()='t']/text()"))
        expected = normalize(base.parent / f"{stem}.pptx.truth.json")
        assert " ".join(expected.text.split()) == " ".join(actual.split())
    workbook = load_workbook(base.with_suffix(".xlsx"), read_only=True)
    try:
        actual_rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
        expected = normalize(base.parent / f"{stem}.xlsx.truth.json")
        assert actual_rows == [
            [cell.text for cell in row] for row in expected.tables[0].rows
        ]
    finally:
        workbook.close()


def test_ssim_detects_content_and_geometry_changes():
    reference = np.full((64, 64), 255, dtype=np.uint8)
    reference[8:32, 8:48] = 0
    changed = reference.copy()
    changed[8:32, 8:48] = 255
    assert page_ssim(reference, reference) == pytest.approx(1)
    assert page_ssim(reference, changed) < 0.9
    assert page_ssim(reference, np.tile(reference, (2, 1))) < 0.95


def test_pdf_ssim_penalizes_missing_pages(tmp_path):
    import pypdfium2 as pdfium

    source = ROOT / "fixtures/heavy/en-map.pdf"
    with pdfium.PdfDocument(source) as document, pdfium.PdfDocument.new() as doubled:
        doubled.import_pages(document)
        doubled.import_pages(document)
        doubled.save(tmp_path / "two-pages.pdf")
    assert pdf_ssim(source, tmp_path / "two-pages.pdf") == 0.5


def test_generated_report_matches_committed_measurements():
    evidence = ROOT / "bench/results/heavy-routes.json"
    capabilities = json.loads(
        (ROOT / "crates/ariad-core/data/capabilities.json").read_text()
    )
    report = render(capabilities, json.loads(evidence.read_text()))
    assert (ROOT / "docs/bench-report.md").read_text() == report


def test_heavy_evidence_covers_required_routes_and_matches_inputs():
    import hashlib

    evidence = json.loads((ROOT / "bench/results/heavy-routes.json").read_text())
    manifest = tomllib.loads((ROOT / "fixtures/manifest.toml").read_text())["fixture"]
    fixtures = {f["id"]: f for f in manifest if "heavy-bench" in f.get("tags", [])}
    assert {f["id"] for f in evidence["fixtures"]} == set(fixtures)
    for record in evidence["fixtures"]:
        fixture = fixtures[record["id"]]
        assert (
            record["sha256"]
            == hashlib.sha256((ROOT / fixture["path"]).read_bytes()).hexdigest()
        )
        assert record["truth"] == fixture["companions"]
        for companion in record["truth"]:
            assert (
                companion["sha256"]
                == hashlib.sha256((ROOT / companion["path"]).read_bytes()).hexdigest()
            )
    samples = evidence["samples"]
    for source in ["doc", "odt", "rtf"]:
        outputs = [
            s
            for s in samples
            if (s["source"], s["target"], s["engine"])
            == (source, "docx", "libreoffice")
        ]
        assert len(outputs) == 2, source
        assert all(s["editability"] > 0.75 for s in outputs)
    for fmt in [
        "docx",
        "pptx",
        "xlsx",
        "odt",
        "ods",
        "odp",
        "doc",
        "rtf",
        "xls",
        "ppt",
    ]:
        outputs = [
            s
            for s in samples
            if (s["source"], s["target"], s["engine"]) == (fmt, "pdf", "libreoffice")
        ]
        assert len(outputs) == 2, fmt
        assert all(
            s["visual_profile"] == "faithful" and 0 <= s["ssim"] <= 1 for s in outputs
        )
    for source in ["pdf", "png"]:
        for ocr in ["tesseract", "rapidocr"]:
            for output in ["markdown", "docx"]:
                routes = [
                    s
                    for s in samples
                    if s["route"]
                    and (s["source"], s["target"], s["ocr"]) == (source, output, ocr)
                ]
                assert len(routes) == 2, (source, output, ocr)
    for source in ["md", "docx"]:
        routes = [
            s
            for s in samples
            if s["route"]
            and (s["source"], s["target"], s["engine"]) == (source, "pdf", "typst")
        ]
        assert len(routes) == 2, source
    assert all(s["wall_ms"] > 0 and s["peak_rss_bytes"] > 0 for s in samples)
    assert all(
        0 <= s["ssim"] <= 1 and s["visual_profile"] == "faithful"
        for s in samples
        if s["route"]
    )
    mapper = [s for s in samples if s["source"] == "docling+json"]
    assert mapper
    known = next(
        s for s in mapper if s["fixture"] == "heavy-en-map-pdf" and s["ocr"] == "off"
    )
    assert known["fidelity"] == 1 and known["editability"] == 1


def test_report_rejects_mixed_measurement_runs():
    capabilities = json.loads(
        (ROOT / "crates/ariad-core/data/capabilities.json").read_text()
    )
    evidence = json.loads((ROOT / "bench/results/heavy-routes.json").read_text())
    evidence["generated_at"] = "2000-01-01T00:00:00Z"
    with pytest.raises(ValueError, match="different runs"):
        render(capabilities, evidence)


def test_native_benchmark_workspace_supports_pandoc_roundtrip(tmp_path):
    from ariad_bench.heavy import Bench, IR
    from ariad_bench.run import find_binary

    ashift = find_binary(
        "ashift",
        None,
        [ROOT / "target" / "release" / "ashift", ROOT / "target" / "debug" / "ashift"],
    )
    pandoc = find_binary("pandoc", None, [ROOT / ".tools/pandoc/bin/pandoc"])
    truth = json.loads((ROOT / "fixtures/heavy/en-map.docx.truth.json").read_text())
    source = tmp_path / "source.json"
    source.write_text(json.dumps(truth["ir"]))
    bench = Bench(ROOT, ashift, pandoc, tmp_path, 1)
    docx, _, _ = bench.engine("pandoc", source, IR, "docx")
    restored, _, _ = bench.engine("pandoc", docx, "docx", IR)
    assert docx.read_bytes().startswith(b"PK")
    assert normalize(restored).text == normalize(truth).text
