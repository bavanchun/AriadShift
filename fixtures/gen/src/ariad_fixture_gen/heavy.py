"""Authored block lists shared by digital, scanned and Office benchmark inputs."""

from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import typst

from ariad_fixture_gen import docx, office, scan
from ariad_fixture_gen.pdf import _content, _source
from ariad_fixture_gen.truth import (
    Cell,
    HeadingBlock,
    ParagraphBlock,
    TableBlock,
    blocks_to_canonical,
    blocks_to_ir,
)


def cases(root: Path) -> list[dict]:
    """Return original source content, never a conversion engine's interpretation."""
    return [
        {
            "id": "en-map",
            "lang": "en",
            "title": "Notes from the Map Room",
            "paragraphs": _source(root, "english-map-room.md")[:2],
            "sections": ["Archive practice", "Reading old maps"],
            "rows": [
                ("Archive item", "Recorded detail"),
                ("Street plan", "Survey date"),
                ("Overlay", "Transparent sheet"),
                ("Map weight", "Brass"),
            ],
        },
        {
            "id": "vi-garden",
            "lang": "vi",
            "title": "Khu vườn đọc sách",
            "paragraphs": _source(root, "vietnamese-garden.md")[:2],
            "sections": ["Lịch chăm vườn", "Ghi chép trong tuần"],
            "rows": [
                ("Điểm dừng", "Phút đi bộ"),
                ("Bến sông", "8"),
                ("Vườn đọc", "14"),
                ("Mái hiên", "6"),
            ],
        },
    ]


def blocks(case: dict) -> list:
    result = [HeadingBlock(1, case["title"])]
    for heading, paragraph in zip(case["sections"], case["paragraphs"], strict=True):
        result.extend([HeadingBlock(2, heading), ParagraphBlock(paragraph)])
    result.append(
        TableBlock(
            [
                [Cell(value, is_header=index == 0) for value in row]
                for index, row in enumerate(case["rows"])
            ]
        )
    )
    return result


def _source_typst(case: dict) -> bytes:
    source = [
        '#set document(author: "AriadShift contributors", date: none)',
        '#set page(paper: "a4", margin: 22mm)',
        f'#set text(font: "Libertinus Serif", size: 12pt, lang: "{case["lang"]}")',
        "#set par(leading: 0.8em)",
        "= " + _content(case["title"]),
        "#columns(2, gutter: 10mm)[",
    ]
    for index, (heading, paragraph) in enumerate(
        zip(case["sections"], case["paragraphs"], strict=True)
    ):
        if index:
            source.append("#colbreak()")
        source.extend(["== " + _content(heading), _content(paragraph), "#parbreak()"])
    source.extend(
        ["]", "#v(6mm)", "#table(columns: (1fr, 1fr), inset: 6pt, stroke: 0.5pt,"]
    )
    for index, row in enumerate(case["rows"]):
        cells = ", ".join("[" + _content(value) + "]" for value in row)
        source.append(("table.header(" + cells + ")," if index == 0 else cells + ","))
    source.append(")")
    return "\n".join(source).encode()


def _office(case: dict) -> tuple[bytes, bytes]:
    presentation = Presentation()
    presentation.core_properties.author = office._AUTHOR
    presentation.core_properties.last_modified_by = office._AUTHOR
    presentation.core_properties.created = office._OOXML_TIME
    presentation.core_properties.modified = office._OOXML_TIME
    presentation.core_properties.title = case["title"]
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = case["title"]
    table = slide.shapes.add_table(
        len(case["rows"]),
        len(case["rows"][0]),
        Inches(1),
        Inches(2),
        Inches(8),
        Inches(3),
    ).table
    for row_index, row in enumerate(case["rows"]):
        for column, value in enumerate(row):
            table.cell(row_index, column).text = value
    deck = BytesIO()
    presentation.save(deck)

    workbook = Workbook()
    workbook.properties.creator = office._AUTHOR
    workbook.properties.lastModifiedBy = office._AUTHOR
    workbook.properties.created = office._OOXML_TIME
    workbook.properties.modified = office._OOXML_TIME
    sheet = workbook.active
    sheet.title = "Records"
    for row in case["rows"]:
        sheet.append(row)
    book = BytesIO()
    workbook.save(book)
    return office._normalized_zip(deck.getvalue()), office._normalized_zip(
        book.getvalue()
    )


def _truth(case: dict, content: list) -> bytes:
    ir = blocks_to_ir(case["title"], case["lang"], content)
    return (
        json.dumps(
            blocks_to_canonical(case["id"], content, ir_doc=ir),
            ensure_ascii=False,
            indent=2,
        )
        + "\n"
    ).encode()


def generate(root: Path) -> dict[str, bytes]:
    outputs = {}
    for case in cases(root):
        stem = f"fixtures/heavy/{case['id']}"
        source = _source_typst(case)
        compile_options = {
            "font_paths": [],
            "ignore_system_fonts": True,
            "timestamp": 1_700_000_000,
        }
        outputs[stem + ".pdf"] = typst.compile(source, format="pdf", **compile_options)
        png = typst.compile(source, format="png", ppi=200, **compile_options)
        degraded = scan._degrade_png(
            png, seed=791, ppi=200, angle=0.2, blur_radius=0.15, jpeg_quality=92
        )
        outputs[stem + "-scan.png"] = degraded
        outputs[stem + "-scan.pdf"] = scan._scan_pdf(degraded, ppi=200)
        content = blocks(case)
        for suffix in [".pdf", "-scan.pdf", "-scan.png", ".md", ".docx"]:
            outputs[stem + suffix + ".truth.json"] = _truth(case, content)
        markdown = ["# " + case["title"]]
        for heading, paragraph in zip(
            case["sections"], case["paragraphs"], strict=True
        ):
            markdown.extend(["## " + heading, paragraph])
        markdown.append(
            "| "
            + " | ".join(case["rows"][0])
            + " |\n| --- | --- |\n"
            + "\n".join("| " + " | ".join(row) + " |" for row in case["rows"][1:])
        )
        outputs[stem + ".md"] = ("\n\n".join(markdown) + "\n").encode()
        outputs[stem + ".docx"] = docx._render_docx(
            docx.DocxSpec(
                path=stem + ".docx",
                doc_id=case["id"],
                title=case["title"],
                lang=case["lang"],
                blocks=content,
            )
        )
        deck, book = _office(case)
        outputs[stem + ".pptx"] = deck
        outputs[stem + ".xlsx"] = book
        outputs[stem + ".pptx.truth.json"] = _truth(case, [content[0], content[-1]])
        outputs[stem + ".xlsx.truth.json"] = _truth(case, [content[-1]])
    return outputs
