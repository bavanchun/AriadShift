"""Generate reproducible OOXML/ODF fixtures and legacy Office fixtures with LibreOffice.

ODT, ODS and ODP files and the dedicated PPTX/XLSX parser fixtures are reproducible.
The Office-normalization samples (DOC, RTF, XLS, XLSX, PPT and PPTX) come from a system
LibreOffice (``soffice``, or ``ASHIFT_SOFFICE``). LibreOffice stamps the conversion time,
so existing converted samples stay as committed; delete one to convert it again.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches, Pt

_FIXED_TIME = (2026, 1, 1, 0, 0, 0)

_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
    'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
    'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
    'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0"'
)

_MEDIA_TYPES = {
    "odt": "application/vnd.oasis.opendocument.text",
    "ods": "application/vnd.oasis.opendocument.spreadsheet",
    "odp": "application/vnd.oasis.opendocument.presentation",
}

# Target extension -> LibreOffice --convert-to argument.
_FILTERS = {
    "doc": "doc:MS Word 97",
    "rtf": "rtf:Rich Text Format",
    "xls": "xls:MS Excel 97",
    "xlsx": "xlsx:Calc MS Excel 2007 XML",
    "ppt": "ppt:MS PowerPoint 97",
    "pptx": "pptx:Impress MS PowerPoint 2007 XML",
}


def _source(root: Path, filename: str) -> list[str]:
    path = root / "fixtures" / "gen" / "src" / "ariad_fixture_gen" / "text" / filename
    return [
        part.strip()
        for part in path.read_text(encoding="utf-8").split("\n\n")
        if part.strip()
    ]


def _zip_entry(
    archive: ZipFile, name: str, data: bytes, *, stored: bool = False
) -> None:
    info = ZipInfo(name, date_time=_FIXED_TIME)
    info.compress_type = ZIP_STORED if stored else ZIP_DEFLATED
    info.external_attr = 0o644 << 16
    archive.writestr(info, data)


def _package(kind: str, title: str, language: str, body: str) -> bytes:
    """Pack a minimal OpenDocument package; the stored ``mimetype`` entry comes first."""
    media_type = _MEDIA_TYPES[kind]
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        'manifest:version="1.3">'
        f'<manifest:file-entry manifest:full-path="/" manifest:version="1.3" manifest:media-type="{media_type}"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>'
        '<manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>\n"
    )
    meta = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-meta {_NS} office:version="1.3"><office:meta>'
        f"<dc:title>{escape(title)}</dc:title><dc:language>{language}</dc:language>"
        "<meta:generator>ariad-fixture-gen</meta:generator>"
        "</office:meta></office:document-meta>\n"
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-content {_NS} office:version="1.3"><office:body>{body}</office:body>'
        "</office:document-content>\n"
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<office:document-styles {_NS} office:version="1.3"><office:styles>'
        f"{_heading_style(1, '18pt')}{_heading_style(2, '14pt')}"
        "</office:styles></office:document-styles>\n"
    )
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        _zip_entry(archive, "mimetype", media_type.encode("ascii"), stored=True)
        _zip_entry(archive, "META-INF/manifest.xml", manifest.encode("utf-8"))
        _zip_entry(archive, "meta.xml", meta.encode("utf-8"))
        _zip_entry(archive, "styles.xml", styles.encode("utf-8"))
        _zip_entry(archive, "content.xml", content.encode("utf-8"))
    return output.getvalue()


def _heading_style(level: int, size: str) -> str:
    """A paragraph style LibreOffice maps to its built-in "Heading N" style."""
    return (
        f'<style:style style:name="Heading_20_{level}" style:display-name="Heading {level}" '
        f'style:family="paragraph" style:default-outline-level="{level}" style:class="text">'
        f'<style:text-properties fo:font-size="{size}" fo:font-weight="bold"/></style:style>'
    )


def _paragraph(text: str) -> str:
    return f"<text:p>{escape(text)}</text:p>"


def _cell(text: str) -> str:
    return f'<table:table-cell office:value-type="string">{_paragraph(text)}</table:table-cell>'


def _number_cell(value: int) -> str:
    return f'<table:table-cell office:value-type="float" office:value="{value}">{_paragraph(str(value))}</table:table-cell>'


def _text_document(paragraphs: list[str]) -> str:
    rows = [("Điểm dừng", "Phút đi bộ"), ("Bến sông", "8"), ("Vườn đọc", "14")]
    parts = [
        "<office:text>",
        '<text:h text:style-name="Heading_20_1" text:outline-level="1">Khu vườn đọc sách</text:h>',
        *(_paragraph(paragraph) for paragraph in paragraphs[:2]),
        '<text:h text:style-name="Heading_20_2" text:outline-level="2">Lịch chăm cây</text:h>',
        "<text:list>",
        *(
            f"<text:list-item>{_paragraph(item)}</text:list-item>"
            for item in (
                "Tưới cây vào sáng thứ bảy",
                "Đổi sách ở bàn gỗ",
                "Ghi chú trên bảng",
            )
        ),
        "</text:list>",
        '<table:table table:name="Lich"><table:table-column table:number-columns-repeated="2"/>',
        *(
            f"<table:table-row>{_cell(left)}{_cell(right)}</table:table-row>"
            for left, right in rows
        ),
        "</table:table>",
        _paragraph(paragraphs[2]),
        "</office:text>",
    ]
    return "".join(parts)


def _spreadsheet() -> str:
    rows = [
        ("Điểm dừng", "Phút đi bộ", "Ghi chú"),
        ("Bến sông", 8, "Có ghế"),
        ("Vườn đọc", 14, "Mái hiên"),
    ]
    body = ['<office:spreadsheet><table:table table:name="Lộ trình">']
    for row in rows:
        cells = "".join(
            _number_cell(value) if isinstance(value, int) else _cell(value)
            for value in row
        )
        body.append(f"<table:table-row>{cells}</table:table-row>")
    body.append("</table:table></office:spreadsheet>")
    return "".join(body)


def _frame(kind: str, y_cm: float, height_cm: float, paragraphs: list[str]) -> str:
    text = "".join(_paragraph(paragraph) for paragraph in paragraphs)
    return (
        f'<draw:frame presentation:class="{kind}" svg:x="2cm" svg:y="{y_cm}cm" '
        f'svg:width="24cm" svg:height="{height_cm}cm"><draw:text-box>{text}</draw:text-box></draw:frame>'
    )


def _presentation(paragraphs: list[str]) -> str:
    slides = [
        (
            "Notes from the Map Room",
            ["Street plans, ferry schedules and hand-drawn routes"],
        ),
        ("Archive practice", [paragraphs[1][:120], "Overlays keep older maps intact"]),
    ]
    body = ["<office:presentation>"]
    for index, (title, lines) in enumerate(slides, start=1):
        body.append(f'<draw:page draw:name="page{index}">')
        body.append(_frame("title", 1.0, 3.0, [title]))
        body.append(_frame("outline", 5.0, 12.0, lines))
        body.append("</draw:page>")
    body.append("</office:presentation>")
    return "".join(body)


def _soffice() -> str | None:
    configured = os.environ.get("ASHIFT_SOFFICE")
    if configured:
        return configured
    return shutil.which("soffice") or shutil.which("libreoffice")


def _convert(soffice: str, source_name: str, source: bytes, target: str) -> bytes:
    """Convert one ODF document with a throwaway LibreOffice profile."""
    with tempfile.TemporaryDirectory(prefix="ariad-fixture-office-") as scratch:
        work = Path(scratch)
        (work / source_name).write_bytes(source)
        profile = (work / "profile").as_uri()
        subprocess.run(
            [
                soffice,
                "--headless",
                "--norestore",
                "--nolockcheck",
                f"-env:UserInstallation={profile}",
                "--convert-to",
                _FILTERS[target],
                "--outdir",
                str(work / "out"),
                str(work / source_name),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
        return (work / "out" / f"{Path(source_name).stem}.{target}").read_bytes()


_OOXML_TIME = datetime(2000, 1, 1)
_AUTHOR = "AriadShift contributors"

_MODIFIED = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _normalized_zip(data: bytes) -> bytes:
    """Rewrites an OOXML package with sorted, stored entries and fixed timestamps."""
    with ZipFile(BytesIO(data)) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    # openpyxl stamps the save time into the core properties.
    parts["docProps/core.xml"] = _MODIFIED.sub(
        rb"\g<1>2000-01-01T00:00:00Z\g<2>", parts["docProps/core.xml"]
    )
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_STORED) as archive:
        for name, content in sorted(parts.items()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_STORED
            info.create_system = 0
            info.external_attr = 0
            archive.writestr(info, content)
    return output.getvalue()


def _slides(root: Path) -> bytes:
    paragraphs = _source(root, "vietnamese-garden.md")
    presentation = Presentation()
    properties = presentation.core_properties
    properties.author = _AUTHOR
    properties.last_modified_by = _AUTHOR
    properties.title = "Khu vườn đọc sách"
    properties.created = _OOXML_TIME
    properties.modified = _OOXML_TIME
    properties.revision = 1

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "Khu vườn đọc sách"
    title_slide.placeholders[1].text = "Một khoảng sân nhỏ sau thư viện phường"

    bullets = presentation.slides.add_slide(presentation.slide_layouts[1])
    bullets.shapes.title.text = "Lịch chăm vườn"
    body = bullets.placeholders[1].text_frame
    body.text = "Thứ bảy: tưới cây và đổi sách"
    for line in (
        "Ngày nắng: học bài dưới giàn hoa giấy",
        "Ngày mưa: chuyển vào mái hiên",
    ):
        body.add_paragraph().text = line

    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "Sách được đổi trong tháng"
    rows = [("Loại sách", "Số quyển"), ("Truyện tranh", "24"), ("Sách nấu ăn", "11")]
    shape = table_slide.shapes.add_table(
        len(rows), 2, Inches(1), Inches(2), Inches(6), Inches(1.5)
    )
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = shape.table.cell(row_index, column_index)
            cell.text = value
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(18)

    prose = presentation.slides.add_slide(presentation.slide_layouts[1])
    prose.shapes.title.text = "Ghi chép"
    prose.placeholders[1].text_frame.text = paragraphs[1]

    output = BytesIO()
    presentation.save(output)
    return _normalized_zip(output.getvalue())


def _workbook() -> bytes:
    workbook = Workbook()
    workbook.properties.creator = _AUTHOR
    workbook.properties.lastModifiedBy = _AUTHOR
    workbook.properties.title = "Repair café log"
    workbook.properties.created = _OOXML_TIME
    workbook.properties.modified = _OOXML_TIME

    log = workbook.active
    log.title = "Repairs"
    rows = [
        ("Item", "Fault", "Minutes", "Fixed"),
        ("Desk lamp", "Loose switch", 15, "yes"),
        ("Toaster", "Broken lever", 40, "yes"),
        ("Radio", "No sound", 55, "no"),
        ("Bicycle light", "Corroded battery contacts", 10, "yes"),
    ]
    for row in rows:
        log.append(row)
    for cell in log[1]:
        cell.font = Font(bold=True)

    volunteers = workbook.create_sheet("Tình nguyện viên")
    volunteers.append(("Tên", "Ca trực"))
    volunteers.append(("Mai", "Sáng thứ bảy"))
    volunteers.append(("Hùng", "Chiều chủ nhật"))
    for cell in volunteers[1]:
        cell.font = Font(bold=True)

    output = BytesIO()
    workbook.save(output)
    return _normalized_zip(output.getvalue())


def generate(root: Path) -> dict[str, bytes]:
    garden = _source(root, "vietnamese-garden.md")
    map_room = _source(root, "english-map-room.md")
    sources = {
        "fixtures/office/vi-garden-notes.odt": _package(
            "odt", "Khu vườn đọc sách", "vi", _text_document(garden)
        ),
        "fixtures/office/vi-route-sheet.ods": _package(
            "ods", "Lộ trình", "vi", _spreadsheet()
        ),
        "fixtures/office/en-map-room-slides.odp": _package(
            "odp", "Notes from the Map Room", "en", _presentation(map_room)
        ),
    }
    derived = {
        "fixtures/office/vi-garden-notes.odt": ("doc", "rtf"),
        "fixtures/office/vi-route-sheet.ods": ("xls", "xlsx"),
        "fixtures/office/en-map-room-slides.odp": ("ppt", "pptx"),
    }

    outputs = dict(sources)
    outputs.update(
        {
            "fixtures/pptx/vi-garden-slides.pptx": _slides(root),
            "fixtures/xlsx/en-repair-log.xlsx": _workbook(),
        }
    )
    soffice = _soffice()
    for source_path, targets in derived.items():
        for target in targets:
            target_path = str(Path(source_path).with_suffix(f".{target}"))
            existing = root / target_path
            if existing.is_file():
                outputs[target_path] = existing.read_bytes()
            elif soffice is not None:
                outputs[target_path] = _convert(
                    soffice, Path(source_path).name, sources[source_path], target
                )
            else:
                raise RuntimeError(
                    f"{target_path} is missing and LibreOffice is not installed; "
                    "install LibreOffice or set ASHIFT_SOFFICE to generate it"
                )
    return outputs
