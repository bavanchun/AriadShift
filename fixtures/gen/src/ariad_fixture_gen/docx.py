"""Generate reproducible DOCX fixtures with structural edge cases."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_STORED, ZipFile, ZipInfo

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

from ariad_fixture_gen.truth import (
    Cell,
    FootnoteBlock,
    HeadingBlock,
    ListBlock,
    PageBreakBlock,
    ParagraphBlock,
    TableBlock,
    UnscoredBlock,
    blocks_to_canonical,
    blocks_to_ir,
)

_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_FOOTNOTE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"


def _source(root: Path, filename: str) -> list[str]:
    path = root / "fixtures" / "gen" / "src" / "ariad_fixture_gen" / "text" / filename
    return [part.strip() for part in path.read_text(encoding="utf-8").split("\n\n") if part.strip()]


def _document(title: str) -> Document:
    document = Document()
    properties = document.core_properties
    properties.author = "AriadShift contributors"
    properties.last_modified_by = "AriadShift contributors"
    properties.title = title
    properties.subject = "Reproducible document-conversion fixture"
    properties.created = datetime(2000, 1, 1)
    properties.modified = datetime(2000, 1, 1)
    return document


def _set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold


def _xml_bytes(root) -> bytes:
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def _add_footnote(parts: dict[str, bytes], text: str) -> None:
    namespaces = {"w": _WORD_NS}
    document_xml = etree.fromstring(parts["word/document.xml"])
    markers = document_xml.xpath(".//w:t[text()='__FOOTNOTE_1__']", namespaces=namespaces)
    if len(markers) != 1:
        raise ValueError("Expected exactly one footnote anchor")

    run = etree.Element(f"{{{_WORD_NS}}}r")
    run_properties = etree.SubElement(run, f"{{{_WORD_NS}}}rPr")
    style = etree.SubElement(run_properties, f"{{{_WORD_NS}}}rStyle")
    style.set(f"{{{_WORD_NS}}}val", "FootnoteReference")
    reference = etree.SubElement(run, f"{{{_WORD_NS}}}footnoteReference")
    reference.set(f"{{{_WORD_NS}}}id", "1")
    marker_run = markers[0].getparent()
    marker_run.getparent().replace(marker_run, run)
    parts["word/document.xml"] = _xml_bytes(document_xml)

    footnotes = etree.Element(f"{{{_WORD_NS}}}footnotes", nsmap={"w": _WORD_NS})
    for footnote_id, separator in ((-1, "separator"), (0, "continuationSeparator")):
        footnote = etree.SubElement(footnotes, f"{{{_WORD_NS}}}footnote")
        footnote.set(f"{{{_WORD_NS}}}type", separator)
        footnote.set(f"{{{_WORD_NS}}}id", str(footnote_id))
        paragraph = etree.SubElement(footnote, f"{{{_WORD_NS}}}p")
        separator_run = etree.SubElement(paragraph, f"{{{_WORD_NS}}}r")
        etree.SubElement(separator_run, f"{{{_WORD_NS}}}{separator}")

    footnote = etree.SubElement(footnotes, f"{{{_WORD_NS}}}footnote")
    footnote.set(f"{{{_WORD_NS}}}id", "1")
    paragraph = etree.SubElement(footnote, f"{{{_WORD_NS}}}p")
    paragraph_properties = etree.SubElement(paragraph, f"{{{_WORD_NS}}}pPr")
    paragraph_style = etree.SubElement(paragraph_properties, f"{{{_WORD_NS}}}pStyle")
    paragraph_style.set(f"{{{_WORD_NS}}}val", "FootnoteText")
    reference_run = etree.SubElement(paragraph, f"{{{_WORD_NS}}}r")
    reference_properties = etree.SubElement(reference_run, f"{{{_WORD_NS}}}rPr")
    reference_style = etree.SubElement(reference_properties, f"{{{_WORD_NS}}}rStyle")
    reference_style.set(f"{{{_WORD_NS}}}val", "FootnoteReference")
    etree.SubElement(reference_run, f"{{{_WORD_NS}}}footnoteRef")
    text_run = etree.SubElement(paragraph, f"{{{_WORD_NS}}}r")
    text_node = etree.SubElement(text_run, f"{{{_WORD_NS}}}t")
    text_node.text = text
    parts["word/footnotes.xml"] = _xml_bytes(footnotes)

    relationships = etree.fromstring(parts["word/_rels/document.xml.rels"])
    used_ids = [int(item.get("Id", "rId0")[3:]) for item in relationships if item.get("Id", "").startswith("rId")]
    relationship = etree.SubElement(relationships, f"{{{_REL_NS}}}Relationship")
    relationship.set("Id", f"rId{max(used_ids, default=0) + 1}")
    relationship.set("Type", _FOOTNOTE_REL)
    relationship.set("Target", "footnotes.xml")
    parts["word/_rels/document.xml.rels"] = _xml_bytes(relationships)

    content_types = etree.fromstring(parts["[Content_Types].xml"])
    override = etree.SubElement(content_types, f"{{{_CONTENT_NS}}}Override")
    override.set("PartName", "/word/footnotes.xml")
    override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
    parts["[Content_Types].xml"] = _xml_bytes(content_types)

    styles = etree.fromstring(parts["word/styles.xml"])
    footnote_text_style = etree.SubElement(styles, f"{{{_WORD_NS}}}style")
    footnote_text_style.set(f"{{{_WORD_NS}}}type", "paragraph")
    footnote_text_style.set(f"{{{_WORD_NS}}}styleId", "FootnoteText")
    footnote_text_name = etree.SubElement(footnote_text_style, f"{{{_WORD_NS}}}name")
    footnote_text_name.set(f"{{{_WORD_NS}}}val", "Footnote Text")
    footnote_reference_style = etree.SubElement(styles, f"{{{_WORD_NS}}}style")
    footnote_reference_style.set(f"{{{_WORD_NS}}}type", "character")
    footnote_reference_style.set(f"{{{_WORD_NS}}}styleId", "FootnoteReference")
    footnote_reference_name = etree.SubElement(footnote_reference_style, f"{{{_WORD_NS}}}name")
    footnote_reference_name.set(f"{{{_WORD_NS}}}val", "Footnote Reference")
    parts["word/styles.xml"] = _xml_bytes(styles)

    settings = etree.fromstring(parts["word/settings.xml"])
    footnote_properties = etree.Element(f"{{{_WORD_NS}}}footnotePr")
    number_format = etree.SubElement(footnote_properties, f"{{{_WORD_NS}}}numFmt")
    number_format.set(f"{{{_WORD_NS}}}val", "decimal")
    compatibility = settings.find(f"{{{_WORD_NS}}}compat")
    if compatibility is None:
        settings.append(footnote_properties)
    else:
        settings.insert(settings.index(compatibility), footnote_properties)
    parts["word/settings.xml"] = _xml_bytes(settings)


def _save(document: Document, *, footnote_text: str | None = None) -> bytes:
    source = BytesIO()
    document.save(source)
    with ZipFile(BytesIO(source.getvalue())) as original:
        parts = {name: original.read(name) for name in original.namelist()}

    if footnote_text is not None:
        _add_footnote(parts, footnote_text)

    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_STORED) as archive:
        for name, content in sorted(parts.items()):
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_STORED
            info.create_system = 0
            info.external_attr = 0
            info.internal_attr = 0
            info.extra = b""
            info.comment = b""
            archive.writestr(info, content)
    return output.getvalue()


@dataclass
class DocxSpec:
    path: str
    doc_id: str
    title: str
    lang: str
    blocks: list[Any]
    margins: tuple[float, float] | None = None
    header_text: str | None = None
    footer_text: str | None = None
    columns: int | None = None
    column_space: int | None = None


def _render_docx(spec: DocxSpec) -> bytes:
    """Render DOCX binary bytes directly from the spec's ordered block list."""
    document = _document(spec.title)
    section = document.sections[0]

    if spec.margins is not None:
        section.top_margin = Inches(spec.margins[0])
        section.bottom_margin = Inches(spec.margins[1])
    if spec.header_text is not None:
        section.header.paragraphs[0].text = spec.header_text
    if spec.footer_text is not None:
        section.footer.paragraphs[0].text = spec.footer_text
    if spec.columns is not None:
        cols_xml = section._sectPr.xpath("./w:cols")
        cols_xml[0].set(qn("w:num"), str(spec.columns))
        cols_xml[0].set(qn("w:space"), str(spec.column_space or 720))

    footnote_text: str | None = None

    for block in spec.blocks:
        if isinstance(block, HeadingBlock):
            document.add_heading(block.text, level=block.level)
        elif isinstance(block, ParagraphBlock):
            if block.quote:
                note = document.add_paragraph(style="Quote")
                note.add_run(block.text).italic = True
            else:
                para = document.add_paragraph(block.text)
                if block.footnote_anchor:
                    para.add_run(block.footnote_anchor)
        elif isinstance(block, ListBlock):
            for item in block.items:
                style = block.style or ("List Number" if block.ordered else "List Bullet")
                document.add_paragraph(item, style=style)
        elif isinstance(block, PageBreakBlock):
            document.add_page_break()
        elif isinstance(block, TableBlock):
            num_cols = sum(cell.colspan for cell in block.rows[0])
            table = document.add_table(rows=len(block.rows), cols=num_cols)
            if block.table_style:
                table.style = block.table_style
            if block.alignment == "CENTER":
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for r_idx, row in enumerate(block.rows):
                c_idx = 0
                for cell in row:
                    target_cell = table.cell(r_idx, c_idx)
                    if cell.colspan > 1:
                        target_cell = target_cell.merge(table.cell(r_idx, c_idx + cell.colspan - 1))
                    if cell.nested_table is not None:
                        nested = target_cell.add_table(rows=1, cols=1)
                        if cell.nested_table.table_style:
                            nested.style = cell.nested_table.table_style
                        _set_cell_text(nested.cell(0, 0), cell.nested_table.rows[0][0].text)
                    else:
                        _set_cell_text(target_cell, cell.text, bold=cell.is_header)
                    c_idx += cell.colspan
        elif isinstance(block, FootnoteBlock):
            footnote_text = block.text

    return _save(document, footnote_text=footnote_text)


def _get_specs(root: Path) -> list[DocxSpec]:
    garden = _source(root, "vietnamese-garden.md")
    river = _source(root, "vietnamese-river.md")
    repair_cafe = _source(root, "english-repair-cafe.md")
    map_room = _source(root, "english-map-room.md")

    return [
        DocxSpec(
            path="fixtures/docx/vi-styled-report.docx",
            doc_id="vi-styled-report",
            title="Khu vườn đọc sách",
            lang="vi",
            margins=(0.8, 0.8),
            header_text="Thư viện phường",
            footer_text="Sổ chăm vườn",
            blocks=[
                HeadingBlock(level=1, text="Khu vườn đọc sách"),
                ParagraphBlock(text=garden[0]),
                ParagraphBlock(text="Mỗi sáng thứ bảy, một nhóm hàng xóm tưới cây và đổi sách.", quote=True),
                HeadingBlock(level=2, text="Sinh hoạt hằng tuần"),
                ParagraphBlock(text=garden[1]),
                ParagraphBlock(text=garden[2]),
            ],
        ),
        DocxSpec(
            path="fixtures/docx/en-numbered-checklist.docx",
            doc_id="en-numbered-checklist",
            title="Repair café opening checklist",
            lang="en",
            blocks=[
                HeadingBlock(level=1, text="Before the café opens"),
                ParagraphBlock(text=repair_cafe[0]),
                ListBlock(items=["Clear the workbench", "Check each lamp", "Label the parts tray"], ordered=True),
                ListBlock(items=["Ask the visitor before replacing a part."], ordered=False),
                ListBlock(items=["Record the repair"], ordered=False, style="List Bullet 2"),
                ParagraphBlock(text=repair_cafe[1]),
            ],
        ),
        DocxSpec(
            path="fixtures/docx/vi-route-table.docx",
            doc_id="vi-route-data",
            title="Bảng hành trình ven sông",
            lang="vi",
            blocks=[
                HeadingBlock(level=1, text="Bảng hành trình ven sông"),
                ParagraphBlock(text=river[0]),
                TableBlock(
                    rows=[
                        [Cell(text="Điểm dừng", is_header=True), Cell(text="Phút đi bộ", is_header=True), Cell(text="Ghi chú", is_header=True)],
                        [Cell(text="Bến sông"), Cell(text="8"), Cell(text="Bóng mát")],
                        [Cell(text="Vườn đọc"), Cell(text="14"), Cell(text="Có ghế dài")],
                        [Cell(text="Lịch khảo sát cập nhật mỗi tháng.", colspan=3)],
                    ],
                    table_style="Light Shading Accent 1",
                    alignment="CENTER",
                ),
                ParagraphBlock(text="Các khoảng thời gian là ước lượng đi bộ."),
            ],
        ),
        DocxSpec(
            path="fixtures/docx/en-footnotes.docx",
            doc_id="en-footnotes-docx",
            title="Map archive note with a footnote",
            lang="en",
            blocks=[
                HeadingBlock(level=1, text="Notes from the Map Room"),
                ParagraphBlock(text=map_room[0], footnote_anchor="__FOOTNOTE_1__"),
                ParagraphBlock(text=map_room[1]),
                FootnoteBlock(text="The archive keeps older map editions beside later revisions."),
            ],
        ),
        DocxSpec(
            path="fixtures/docx/edge-case-merged-table.docx",
            doc_id="synthetic-merged-table",
            title="Synthetic merged-table edge case",
            lang="en",
            blocks=[
                HeadingBlock(level=1, text="Sparse route record"),
                TableBlock(
                    rows=[
                        [Cell(text="Route / Tuyến", is_header=True, colspan=4)],
                        [Cell(text="Riverbank"), Cell(text="Bến sông"), Cell(text=""), Cell(text="8 min")],
                        [
                            Cell(text="Nested record", colspan=2),
                            Cell(text="", nested_table=TableBlock(rows=[[Cell(text="empty neighbor")]], table_style="Table Grid")),
                            Cell(text=""),
                        ],
                    ],
                    table_style="Table Grid",
                ),
                ParagraphBlock(text=""),
                UnscoredBlock(reason="nested table inside cell cannot be expressed in canonical form"),
            ],
        ),
        DocxSpec(
            path="fixtures/docx/edge-case-two-column.docx",
            doc_id="synthetic-two-column",
            title="Synthetic two-column reading-order edge case",
            lang="en",
            columns=2,
            column_space=720,
            blocks=[
                HeadingBlock(level=1, text="Two-column archive note"),
                ParagraphBlock(text=map_room[0]),
                ParagraphBlock(text=map_room[1]),
                PageBreakBlock(),
                ParagraphBlock(text="This paragraph follows the page break and preserves document order."),
            ],
        ),
    ]


def generate(root: Path) -> dict[str, bytes]:
    """Generate all DOCX fixtures directly from their ordered block specifications."""
    specs = _get_specs(root)
    return {spec.path: _render_docx(spec) for spec in specs}


def generate_truth(root: Path) -> dict[str, dict[str, Any]]:
    """Build canonical truth companions directly from the ordered block specifications."""
    specs = _get_specs(root)
    truth_dict: dict[str, dict[str, Any]] = {}
    companion_name_map = {
        "vi-styled-report": "fixtures/docx/vi-styled-report.truth.json",
        "en-numbered-checklist": "fixtures/docx/en-numbered-checklist.truth.json",
        "vi-route-data": "fixtures/docx/vi-route-data.truth.json",
        "en-footnotes-docx": "fixtures/docx/en-footnotes-docx.truth.json",
        "synthetic-merged-table": "fixtures/docx/synthetic-merged-table.truth.json",
        "synthetic-two-column": "fixtures/docx/synthetic-two-column.truth.json",
    }
    for spec in specs:
        truth_path = companion_name_map[spec.doc_id]
        ir_doc = blocks_to_ir(spec.title, spec.lang, spec.blocks)
        truth_dict[truth_path] = blocks_to_canonical(spec.doc_id, spec.blocks, ir_doc=ir_doc)
    return truth_dict
